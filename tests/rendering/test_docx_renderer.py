import pytest
from docx import Document

from core.rendering.docx_renderer import render_docx


@pytest.mark.rendering
def test_render_docx_creates_file(tmp_path, sample_resume_content):
    output_path = tmp_path / "resume.docx"

    result_path = render_docx(sample_resume_content, str(output_path))

    assert output_path.exists()
    assert result_path == str(output_path)


@pytest.mark.rendering
def test_render_docx_contains_expected_text(tmp_path, sample_resume_content):
    output_path = tmp_path / "resume.docx"

    render_docx(sample_resume_content, str(output_path))

    document = Document(str(output_path))
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "SAMPLE NAME | Software Engineer" in full_text
    assert "Example Corp - Software Engineer" in full_text
    assert "Sample project" in full_text
    assert "Sample bullet one" in full_text
    assert "SKILLS" in full_text


@pytest.mark.rendering
def test_render_docx_respects_custom_class_style(tmp_path, sample_resume_content):
    output_path = tmp_path / "resume.docx"

    render_docx(
        sample_resume_content,
        str(output_path),
        style={"classes": {"name": {"size": 20}}},
    )

    document = Document(str(output_path))
    name_run = document.paragraphs[0].runs[0]

    assert name_run.font.size.pt == 20


@pytest.mark.rendering
def test_render_docx_respects_path_override(tmp_path, sample_resume_content):
    output_path = tmp_path / "resume.docx"

    render_docx(
        sample_resume_content,
        str(output_path),
        style={"overrides": {"experience[0].company_role": {"bold": False}}},
    )

    document = Document(str(output_path))
    company_role_paragraph = next(p for p in document.paragraphs if "Example Corp" in p.text)

    assert company_role_paragraph.runs[0].bold is False