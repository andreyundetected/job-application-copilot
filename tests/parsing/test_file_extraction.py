import pytest
from docx import Document

from core.parsing.file_extraction import extract_text


@pytest.mark.parsing
def test_extract_text_from_txt(tmp_path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text("Sample text content", encoding="utf-8")

    result = extract_text(str(file_path))

    assert result == "Sample text content"


@pytest.mark.parsing
def test_extract_text_from_docx(tmp_path):
    file_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Sample paragraph one")
    document.add_paragraph("Sample paragraph two")
    document.save(str(file_path))

    result = extract_text(str(file_path))

    assert "Sample paragraph one" in result
    assert "Sample paragraph two" in result


@pytest.mark.parsing
def test_extract_text_unsupported_extension(tmp_path):
    file_path = tmp_path / "sample.xyz"
    file_path.write_text("content", encoding="utf-8")

    with pytest.raises(ValueError):
        extract_text(str(file_path))