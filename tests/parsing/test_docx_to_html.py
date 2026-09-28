import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from core.parsing.docx_to_html import docx_to_html


@pytest.mark.parsing
def test_docx_to_html_keeps_run_formatting_and_lists(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Sample Name")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    document.add_paragraph("First bullet", style="List Bullet")
    document.add_paragraph("Second bullet", style="List Bullet")
    path = tmp_path / "sample.docx"
    document.save(str(path))

    html = docx_to_html(str(path))

    assert "font-size:16pt" in html
    assert "font-weight:bold" in html
    assert "color:#595959" in html
    assert html.count("<li>") == 2
    assert html.count("data-eid=") == 2


@pytest.mark.parsing
def test_docx_to_html_skips_empty_paragraphs(tmp_path):
    document = Document()
    document.add_paragraph("")
    document.add_paragraph("Only text")
    path = tmp_path / "sample.docx"
    document.save(str(path))

    html = docx_to_html(str(path))

    assert html.count("<p") == 1