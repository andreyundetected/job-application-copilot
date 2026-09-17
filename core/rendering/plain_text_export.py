from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
import html as html_module


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in text.split("\n\n") if p.strip()]


def render_plain_text_to_docx(text: str, output_path: str) -> str:
    document = Document()
    for paragraph_text in _paragraphs(text):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(paragraph_text)
        run.font.size = Pt(11)
        run.font.name = "Calibri"
    document.save(output_path)
    return output_path


def render_plain_text_to_pdf(text: str, output_path: str) -> str:
    document = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=0.9 * inch,
        rightMargin=0.9 * inch,
        topMargin=0.8 * inch,
        bottomMargin=0.8 * inch,
    )
    style = ParagraphStyle("body", fontName="Helvetica", fontSize=11, leading=15, spaceAfter=12)

    elements = []
    for paragraph_text in _paragraphs(text):
        elements.append(Paragraph(html_module.escape(paragraph_text).replace("\n", "<br/>"), style))
        elements.append(Spacer(1, 4))

    document.build(elements)
    return output_path