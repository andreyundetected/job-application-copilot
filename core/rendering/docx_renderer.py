from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from core.rendering.base import color_hex, merge_style, resolve_style
from core.rendering.html_export import parse_resume_html


def _rgb(color_name: str) -> RGBColor:
    hex_value = color_hex(color_name).lstrip("#")
    return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))


def _apply_run_style(run, resolved: dict, font_name: str):
    run.font.size = Pt(resolved.get("size", 10))
    run.font.name = font_name
    run.bold = resolved.get("bold", False)
    run.italic = resolved.get("italic", False)
    run.font.color.rgb = _rgb(resolved.get("color", "black"))


def _add_styled_paragraph(
    document, text, class_name, style, path=None, alignment=None, bullet=False
):
    resolved = resolve_style(style, class_name, path)
    paragraph = document.add_paragraph(style="List Bullet" if bullet else None)
    if alignment is not None:
        paragraph.alignment = alignment
    run = paragraph.add_run(text)
    _apply_run_style(run, resolved, style["font_name"])
    return paragraph


def _add_content_blocks(document, blocks, path_prefix, style):
    for index, block in enumerate(blocks or []):
        block_path = f"{path_prefix}[{index}]"
        block_type = block.get("type")

        if block_type == "heading":
            _add_styled_paragraph(document, block.get("text", ""), "heading", style, path=block_path)
        elif block_type == "bullet_list":
            for item in block.get("items", []):
                _add_styled_paragraph(document, item, "bullet", style, path=block_path, bullet=True)
        else:
            _add_styled_paragraph(document, block.get("text", ""), "body", style, path=block_path)


def _add_contacts_paragraph(document, contacts, style):
    resolved = resolve_style(style, "contacts")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if style["contacts_layout"] == "inline":
        run = paragraph.add_run(" | ".join(contacts))
        _apply_run_style(run, resolved, style["font_name"])
    else:
        for index, contact in enumerate(contacts):
            run = paragraph.add_run(contact)
            _apply_run_style(run, resolved, style["font_name"])
            if index < len(contacts) - 1:
                paragraph.add_run().add_break()


def _add_section_header(document, text, style):
    _add_styled_paragraph(document, text, "section_header", style)


def _add_experience_entry(document, entry, index, style):
    path_prefix = f"experience[{index}]"

    company_role_text = f"{entry['company']} - {entry['role']}"
    _add_styled_paragraph(
        document, company_role_text, "company_role", style, path=f"{path_prefix}.company_role"
    )

    meta_text = f"{entry['location']} - {entry['dates']}"
    _add_styled_paragraph(document, meta_text, "meta", style, path=f"{path_prefix}.meta")

    if entry.get("employment_type"):
        _add_styled_paragraph(
            document,
            entry["employment_type"],
            "employment_type",
            style,
            path=f"{path_prefix}.employment_type",
        )

    _add_content_blocks(document, entry.get("content", []), f"{path_prefix}.content", style)


DEFAULT_SECTION_LABELS = {
    "summary": "SUMMARY",
    "experience": "EXPERIENCE",
    "skills": "SKILLS",
}


def render_docx(content: dict, output_path: str, style: dict | None = None) -> str:
    style = merge_style(style)
    document = Document()

    section_labels = {**DEFAULT_SECTION_LABELS, **(content.get("section_labels") or {})}

    _add_styled_paragraph(
        document, content["name"], "name", style, alignment=WD_ALIGN_PARAGRAPH.CENTER
    )

    if content.get("contacts"):
        _add_contacts_paragraph(document, content["contacts"], style)

    if content.get("summary"):
        _add_section_header(document, section_labels["summary"], style)
        _add_styled_paragraph(document, content["summary"], "body", style, path="summary")

    if content.get("experience"):
        _add_section_header(document, section_labels["experience"], style)
        for index, entry in enumerate(content["experience"]):
            _add_experience_entry(document, entry, index, style)

    for index, section in enumerate(content.get("extra_sections", [])):
        path_prefix = f"extra_sections[{index}]"
        _add_styled_paragraph(document, section["heading"], "extra_heading", style, path=f"{path_prefix}.heading")
        _add_content_blocks(document, section.get("content", []), f"{path_prefix}.content", style)

    if content.get("skills"):
        _add_section_header(document, section_labels["skills"], style)
        for skill_line in content["skills"]:
            text = f"{skill_line['label']}: {', '.join(skill_line['items'])}"
            _add_styled_paragraph(document, text, "body", style)

    document.save(output_path)
    return output_path


def render_html_export_to_docx(html_content: str, output_path: str, font_name: str = "Calibri") -> str:
    blocks = parse_resume_html(html_content)
    document = Document()

    for block in blocks:
        paragraph = document.add_paragraph(style="List Bullet" if block["type"] == "bullet" else None)
        if block["align"] == "center":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block["align"] == "right":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        for run_data in block["runs"]:
            if run_data["text"] == "\n":
                paragraph.add_run().add_break()
                continue

            run = paragraph.add_run(run_data["text"])
            run.font.name = font_name
            run.font.size = Pt(run_data.get("size") or 10)
            run.bold = bool(run_data.get("bold"))
            run.italic = bool(run_data.get("italic"))

            color = run_data.get("color")
            if color:
                hex_value = color.lstrip("#")
                if len(hex_value) == 6:
                    run.font.color.rgb = RGBColor(
                        int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16)
                    )

    document.save(output_path)
    return output_path